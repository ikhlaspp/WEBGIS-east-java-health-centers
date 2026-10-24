"""Freeze an explicit source inventory; availability does not certify data quality."""

import hashlib
import io
import json
from decimal import Decimal
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

from .facilities import completeness, parse_csv, parse_pdf
from .http_client import atomic_json
from .indicators import parse_workbook, summarize
from .sources import snapshot


def local_path(root, value, within=None):
    path = (Path(root) / value.replace('\\', '/')).resolve()
    boundary = (Path(root) / within).resolve() if within else Path(root).resolve()
    if not path.is_relative_to(boundary):
        raise ValueError('Source path escapes its permitted directory')
    return path


def checked_bytes(path, expected):
    data = path.read_bytes()
    if len(data) != expected['bytes'] or hashlib.sha256(data).hexdigest() != expected['sha256']:
        raise ValueError('Source bytes differ from the approved inventory')
    return data


def freeze(root, plan, output):
    """Fail closed on inventory drift. Only publish a manifest after all checks pass."""
    root = Path(root).resolve()
    output = local_path(root, str(output), 'data/raw')
    if output.suffix != '.json':
        raise ValueError('Acquisition manifest must be a JSON file, not a raw snapshot')
    baseline_path = local_path(root, plan['archive_inventory'])
    baseline = json.loads(checked_bytes(baseline_path, plan['archive_inventory_identity']))['files']
    expected_paths = [item['path'] for item in baseline]
    actual_paths = {p.relative_to(root).as_posix() for p in (root / 'Archives').rglob('*') if p.is_file()}
    if len(set(expected_paths)) != len(expected_paths) or set(expected_paths) != actual_paths:
        raise ValueError('Archive inventory has missing, unexpected or duplicate paths')
    if len(baseline) != plan['expected_archive_files']:
        raise ValueError('Archive count differs from source plan')
    if len(plan['external_receipts']) != plan['expected_external_files']:
        raise ValueError('External count differs from source plan')
    config_evidence = []
    for item in plan.get('configuration', []):
        checked_bytes(local_path(root, item['path']), item)
        config_evidence.append(item)
    entries = []
    # Verify every input before writing snapshots; a partial run never publishes a manifest.
    for item in baseline:
        checked_bytes(local_path(root, item['path'], 'Archives'), item)
    identities = set()
    for receipt in plan['external_receipts']:
        identity = (receipt['source_id'], receipt['edition'], receipt['source_url'])
        if identity in identities:
            raise ValueError('Duplicate external source identity')
        identities.add(identity)
        path = local_path(root, receipt['raw_path'], 'data/raw')
        if path == output:
            raise ValueError('Manifest cannot overwrite a source')
        checked_bytes(path, receipt)
        entries.append({**receipt, 'raw_path': path.relative_to(root).as_posix(),
                        'role': 'selected_official_source', 'data_quality_accepted': False})
    for item in baseline:
        path = local_path(root, item['path'], 'Archives')
        receipt = snapshot(checked_bytes(path, item), root / 'data/raw',
                           'archive-' + hashlib.sha256(item['path'].encode()).hexdigest()[:16],
                           'archive-baseline', item['path'], 'existing_archive')
        receipt['raw_path'] = Path(receipt['raw_path']).relative_to(root).as_posix()
        receipt.update(original_path=item['path'], original_format=path.suffix.lower(),
                       role='supporting_context' if path.suffix.lower() in ('.ipynb', '.docx') else 'historical_source',
                       observation_period=None, period_status='requires_table_level_evidence',
                       data_quality_accepted=False)
        entries.append(receipt)
    manifest = {'schema_version': 1, 'availability_complete': True, 'data_quality_accepted': False,
                'analysis_mode': 'mixed_period_exploration', 'failures': [],
                'archive_files': len(baseline), 'external_files': len(plan['external_receipts']),
                'files': entries, 'configuration': config_evidence,
                'plan_sha256': hashlib.sha256(json.dumps(plan, sort_keys=True).encode()).hexdigest()}
    atomic_json(output, manifest)
    return manifest


def reconcile(root, manifest):
    """Run adapters on frozen bytes and retain discrepancies without fuzzy joins."""
    root = Path(root).resolve()
    originals = {item['original_path']: item for item in manifest['files'] if 'original_path' in item}

    def source_path(receipt):
        path = local_path(root, receipt['raw_path'], 'data/raw')
        checked_bytes(path, receipt)
        return path

    archived_receipt = {**originals['Archives/Puskesmas_2024.pdf'], 'edition': '2024-S1'}
    official_receipt, = [r for r in manifest['files'] if r['source_id'] == 'kemkes-facilities']
    archive = parse_pdf(source_path(archived_receipt), archived_receipt)
    official = parse_pdf(source_path(official_receipt), official_receipt)
    csv_receipt = originals['Archives/puskesmas_jatim_alamat.csv']
    csv_result = parse_csv(source_path(csv_receipt), csv_receipt)
    fields = lambda result: {r['fields']['kode_puskesmas']: r['fields'] for r in result['rows']}
    archive_fields, official_fields, csv_fields = map(fields, (archive, official, csv_result))
    facilities = {'archive': completeness(archive), 'official': completeness(official),
                  'csv': completeness(csv_result),
                  'pdf_field_differences': sum(archive_fields.get(k) != official_fields.get(k)
                                               for k in archive_fields.keys() | official_fields.keys()),
                  'csv_codes_match': csv_fields.keys() == official_fields.keys(),
                  'archive_pages': archive['pages_scanned'], 'official_pages': official['pages_scanned']}
    if not all(facilities[k]['complete'] for k in ('archive', 'official', 'csv')):
        raise ValueError('Facility acquisition does not reconcile to required coverage')
    combined_receipt = originals['Archives/Data_Gabungan_Kecamatan_Jawa_Timur.xlsx']
    book = load_workbook(io.BytesIO(source_path(combined_receipt).read_bytes()), read_only=True, data_only=False)
    combined = [row for row in book['Gabungan'].iter_rows(min_row=3, values_only=True)
                if row and isinstance(row[0], (int, float)) and not isinstance(row[0], bool)]
    book.close()
    key = lambda regency, district: (' '.join(regency.split()).casefold(), ' '.join(district.split()).casefold())
    combined_by_key = {key(row[1], row[2]): row for row in combined}
    if len(combined_by_key) != len(combined):
        raise ValueError('Duplicate exact historical comparison keys')
    indicators = {}
    for indicator, token, column in [('population', 'Data_Penduduk_', 3), ('area', 'Data_Luas_', 4)]:
        receipt, = [r for p, r in originals.items() if token in p]
        parsed = parse_workbook(source_path(receipt), receipt, indicator)
        rows = {key(row['regency_raw'], row['district_raw']): row for row in parsed['rows']}
        if len(rows) != len(parsed['rows']) or parsed['diagnostics']:
            raise ValueError('Indicator adapter has duplicate keys or diagnostics')
        differences = []
        for k in sorted(rows.keys() & combined_by_key.keys()):
            current, historical = rows[k]['value'], combined_by_key[k][column]
            if current is None or historical is None or Decimal(str(current)) != Decimal(str(historical)):
                differences.append({'regency': k[0], 'district': k[1], 'source_value': current,
                                    'combined_value': historical, 'locator': rows[k]['locator']})
        indicators[indicator] = {'report': summarize(parsed), 'combined_rows': len(combined),
                                 'source_missing_in_combined': sorted(rows.keys() - combined_by_key.keys()),
                                 'combined_extra_keys': sorted(combined_by_key.keys() - rows.keys()),
                                 'value_differences': differences, 'source_hash': receipt['sha256'],
                                 'combined_hash': combined_receipt['sha256']}
    pages = []
    for path, receipt in originals.items():
        if path.startswith('Archives/per kecamatan/') and path.endswith('.pdf'):
            with pdfplumber.open(source_path(receipt)) as pdf:
                pages.append({'path': path, 'sha256': receipt['sha256'], 'pages': len(pdf.pages),
                              'observation_period': None, 'table_extraction_pending': True})
    return {'facilities': facilities, 'indicators': indicators, 'archived_bps_pdfs': pages,
            'data_quality_accepted': False,
            'comparison_rule': 'Exact casefolded names with whitespace collapsed; no fuzzy aliases or value correction.'}
