"""Parse source indicator values with explicit units, years and review states."""

import csv
import hashlib
import json
import re
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PARSER_VERSION = 'indicators-1.0'
AREA_EXCEPTIONS = {'Malang': 2019, 'Tulungagung': 2021, 'Kota Probolinggo': 2021,
                   'Lumajang': 2025, 'Kota Madiun': 2025}
OCR_REGIONS = {'Bojonegoro', 'Jember', 'Lamongan', 'Lumajang', 'Kota Madiun'}


def parse_number(raw, locale):
    if raw is None or (isinstance(raw, str) and raw.strip() in ('', '-', '...', '\u2026')):
        return None
    if locale not in ('id_ID', 'en_US'):
        raise ValueError('An explicit supported source locale is required')
    if isinstance(raw, bool):
        raise ValueError('Boolean is not an indicator value')
    if isinstance(raw, (int, float, Decimal)):
        value = Decimal(str(raw))
    else:
        text = str(raw).strip().replace('\u00a0', ' ')
        if text.startswith('='):
            raise ValueError('Formula requires an independently verified value')
        thousands, decimal = ('.', ',') if locale == 'id_ID' else (',', '.')
        pattern = rf'[+-]?(?:\d{{1,3}}(?:{re.escape(thousands)}\d{{3}})+|\d+)(?:{re.escape(decimal)}\d+)?'
        if not re.fullmatch(pattern, text):
            raise ValueError('Number does not match the declared source locale')
        try:
            value = Decimal(text.replace(thousands, '').replace(decimal, '.'))
        except InvalidOperation as exc:
            raise ValueError('Invalid decimal') from exc
    if not value.is_finite() or value < 0:
        raise ValueError('Indicators must be finite and non-negative')
    return value


def convert_value(raw, indicator, unit, locale):
    factors = {'population': {'jiwa': Decimal(1), 'ribu jiwa': Decimal(1000)},
               'area': {'km2': Decimal(1), 'ha': Decimal('0.01'), 'hektar': Decimal('0.01')}}
    if indicator not in factors or unit not in factors[indicator]:
        raise ValueError('Unsupported indicator/unit combination')
    value = parse_number(raw, locale)
    if value is None:
        return None
    value *= factors[indicator][unit]
    if indicator == 'population':
        if value != value.to_integral_value():
            raise ValueError('Population conversion must yield whole people')
        return int(value)
    return format(value, 'f')


def record_value(raw, district, metadata, locator, raw_cells=None):
    row = {**metadata, 'district_raw': district, 'locator': locator, 'raw_value': raw,
           'raw_cells': raw_cells, 'parser_version': PARSER_VERSION}
    try:
        year = metadata.get('observation_year')
        if not isinstance(year, int) or not 1900 <= year <= 2100 or not metadata.get('year_evidence'):
            raise ValueError('Observation year/evidence is missing')
        if not metadata.get('unit_evidence'):
            raise ValueError('Unit evidence is missing')
        row['value'] = convert_value(raw, metadata['indicator'], metadata['unit'], metadata['locale'])
        row['status'] = 'missing' if row['value'] is None else 'staging_unverified'
    except (ValueError, KeyError) as exc:
        row.update(value=None, status='quarantined', reason=str(exc))
    return row


def parse_workbook(path, manifest, indicator):
    """Read the archived derivative layout; its notes are claims, not BPS verification."""
    path = Path(path)
    if hashlib.sha256(path.read_bytes()).hexdigest() != manifest['sha256']:
        raise ValueError('Source hash mismatch')
    result = {'rows': [], 'aggregates': [], 'diagnostics': [], 'notes': []}
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'Catatan' not in workbook.sheetnames:
            raise ValueError('Workbook lacks its source metadata sheet')
        notes = [{'locator': f'xlsx/Catatan/{cell.coordinate}', 'text': str(cell.value)}
                 for row in workbook['Catatan'] for cell in row if cell.value is not None]
        result['notes'] = notes
        combined = '\n'.join(note['text'] for note in notes)
        years = re.search(r'data tahun\s+(\d{4})', combined, flags=re.I)
        publication = re.search(r'Dalam Angka\s+(\d{4})', combined)
        default_year = int(years.group(1)) if years else None
        for sheet in workbook:
            if sheet.title == 'Catatan':
                continue
            rows = list(sheet.iter_rows())
            expected = ('Kecamatan', 'Laki-laki', 'Perempuan', 'Jumlah') if indicator == 'population' else ('Kecamatan', 'Luas (km2)')
            header = tuple(cell.value for cell in rows[1]) if len(rows) > 1 else ()
            if header != expected:
                result['diagnostics'].append({'code': 'header_mismatch', 'sheet': sheet.title,
                                              'raw_header': header})
                continue
            year = default_year
            if indicator == 'area' and sheet.title in AREA_EXCEPTIONS:
                year = AREA_EXCEPTIONS[sheet.title]
                if not re.search(re.escape(sheet.title) + rf'\s+{year}\b', combined):
                    year = None
            unit = 'jiwa' if indicator == 'population' else 'km2'
            metadata = {'source_id': manifest['source_id'], 'source_hash': manifest['sha256'],
                        'source_url': manifest['source_url'], 'acquired_at': manifest['acquired_at'],
                        'source_kind': 'archived_derivative_workbook', 'regency_raw': sheet.title,
                        'indicator': indicator, 'unit': unit, 'locale': 'id_ID',
                        'observation_year': year, 'publication_year': int(publication.group(1)) if publication else None,
                        'year_evidence': notes, 'unit_evidence': {'header': header, 'notes': notes},
                        'review_required': True, 'review_reason': 'Verify derivative against original BPS source',
                        'origin_claim': 'OCR transcription in workbook notes' if indicator == 'population' and sheet.title in OCR_REGIONS else 'Workbook source notes'}
            if indicator == 'population' and 'Jumlah Penduduk' not in combined:
                metadata['unit_evidence'] = None
            index = 3 if indicator == 'population' else 1
            for row_number, cells in enumerate(rows[2:], 3):
                district = cells[0].value
                if district is None and all(cell.value is None for cell in cells):
                    continue
                if isinstance(district, str) and district.startswith('Sumber ') and cells[index].value is None:
                    result['notes'].append({'locator': f'xlsx/{sheet.title}/A{row_number}', 'text': district})
                    continue
                row = record_value(cells[index].value, district, metadata,
                                   f'xlsx/{sheet.title}/{get_column_letter(index + 1)}{row_number}',
                                   [cell.value for cell in cells])
                if district is None:
                    row.update(status='quarantined', reason='District name missing')
                result['aggregates' if str(district).upper() == 'TOTAL' else 'rows'].append(row)
    finally:
        workbook.close()
    return result


def parse_csv_table(path, manifest, spec):
    if hashlib.sha256(Path(path).read_bytes()).hexdigest() != manifest['sha256']:
        raise ValueError('Source hash mismatch')
    result = {'rows': [], 'aggregates': [], 'diagnostics': []}
    with Path(path).open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle, delimiter=spec.get('delimiter', ','))
        required = {spec['district_column'], spec['value_column']}
        if not required.issubset(reader.fieldnames or []):
            return {**result, 'diagnostics': [{'code': 'header_mismatch', 'raw_header': reader.fieldnames}]}
        for number, raw in enumerate(reader, 2):
            metadata = {**spec, 'source_id': manifest['source_id'], 'source_hash': manifest['sha256'],
                        'source_url': manifest['source_url'], 'acquired_at': manifest['acquired_at']}
            district = raw[spec['district_column']]
            row = record_value(raw[spec['value_column']], district, metadata, f'csv/row/{number}', raw)
            result['aggregates' if district in spec.get('total_labels', ['TOTAL']) else 'rows'].append(row)
    return result


def angled_glyph(obj):
    matrix = obj.get('matrix', (1, 0, 0, 1, 0, 0))
    return obj.get('object_type') == 'char' and abs(matrix[1]) > 0.1 and abs(matrix[2]) > 0.1


def save_ocr_review(page, manifest, output_directory, engine=None):
    """Keep OCR as unverified evidence, never silently promote recognized numbers."""
    import numpy as np

    if engine is None:
        from rapidocr_onnxruntime import RapidOCR
        engine = RapidOCR()
    folder = Path(output_directory)
    folder.mkdir(parents=True, exist_ok=True)
    name = f'{manifest["sha256"]}-page-{page.page_number}'
    image_path = folder / (name + '.png')
    rendered = page.to_image(resolution=150).original.convert('RGB')
    rendered.save(image_path)
    tokens, timing = engine(np.asarray(rendered))
    evidence = {'source_hash': manifest['sha256'], 'source_id': manifest['source_id'],
                'pdf_page': page.page_number, 'image_path': str(image_path),
                'image_sha256': hashlib.sha256(image_path.read_bytes()).hexdigest(),
                'engine': 'rapidocr_onnxruntime', 'status': 'requires_visual_review',
                'tokens': [{'box': box, 'text': text, 'confidence': float(score)}
                           for box, text, score in (tokens or [])]}
    evidence_path = folder / (name + '.json')
    evidence_path.write_text(json.dumps(evidence, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    return str(evidence_path)


def parse_pdf_table(path, manifest, spec, ocr_directory=None, ocr_engine=None):
    """Parse a visually selected BPS table; refuse unverified metadata/changed headers."""
    if hashlib.sha256(Path(path).read_bytes()).hexdigest() != manifest['sha256']:
        raise ValueError('Source hash mismatch')
    result = {'rows': [], 'aggregates': [], 'diagnostics': []}
    with pdfplumber.open(path) as pdf:
        page = pdf.pages[spec['pdf_page'] - 1]
        chars = [char for char in page.chars if angled_glyph(char)]
        watermark = ''.join(char['text'] for char in chars)
        filtered = page
        if re.fullmatch(r'https?://[a-z0-9.-]+\.bps\.go\.id/?', watermark):
            filtered = page.filter(lambda obj: not angled_glyph(obj))
            result['extraction_note'] = {'operation': 'exclude_rotated_BPS_URL_watermark', 'text': watermark}
        text = filtered.extract_text() or ''
        if not text.strip():
            diagnostic = {'code': 'requires_ocr', 'page': spec['pdf_page']}
            if ocr_directory is not None:
                diagnostic.update(code='ocr_review_required',
                                  evidence_path=save_ocr_review(page, manifest, ocr_directory, ocr_engine))
            result['diagnostics'].append(diagnostic)
            return result
        normalized_text = ' '.join(text.split())
        if any(' '.join(spec[key].split()) not in normalized_text for key in ('year_evidence', 'table_id')):
            result['diagnostics'].append({'code': 'metadata_mismatch', 'page': spec['pdf_page']})
            return result
        tables = filtered.extract_tables()
        if len(tables) <= spec.get('table_index', 0):
            result['diagnostics'].append({'code': 'table_not_detected', 'page': spec['pdf_page'],
                                          'raw_text': text})
            return result
        table = tables[spec.get('table_index', 0)]
        value_column = spec['value_column']
        header = ' '.join(str(table[0][value_column] or '').split())
        if spec['unit_evidence'] not in header:
            result['diagnostics'].append({'code': 'header_mismatch', 'raw_header': table[0]})
            return result
        for index, cells in enumerate(table[spec['data_start_row'] - 1:], spec['data_start_row']):
            metadata = {**spec, 'source_id': manifest['source_id'], 'source_hash': manifest['sha256'],
                        'source_url': manifest['source_url'], 'acquired_at': manifest['acquired_at'],
                        'source_kind': 'bps_pdf_table'}
            district = ' '.join(str(cells[spec['district_column']] or '').split())
            row = record_value(cells[value_column], district, metadata,
                               f'pdf/page/{spec["pdf_page"]}/table/{spec["table_id"]}/row/{index}', cells)
            result['aggregates' if district in spec['total_labels'] else 'rows'].append(row)
    return result


def summarize(result):
    return {'rows': len(result['rows']), 'aggregates': len(result['aggregates']),
            'diagnostics': len(result['diagnostics']),
            'statuses': dict(Counter(row['status'] for row in result['rows'])),
            'observation_years': dict(Counter(str(row.get('observation_year')) for row in result['rows']))}
