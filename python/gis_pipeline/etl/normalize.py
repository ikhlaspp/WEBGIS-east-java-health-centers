"""Deterministic raw-preserving normalization; no mapping or quality promotion."""

import csv
import hashlib
import io
import json
import re
import unicodedata
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .acquire import checked_bytes, local_path
from .facilities import FIELDS, parse_pdf
from .http_client import atomic_json
from .indicators import convert_value, parse_workbook

VERSION = 'normalize-1.0'
COORDINATE_HEADERS = (*FIELDS, 'query_alamat', 'latitude', 'longitude', 'status_geocoding')


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(',', ':'), allow_nan=False).encode()).hexdigest()


def text(value):
    if not isinstance(value, str):
        raise ValueError('EXPECTED_TEXT')
    return unicodedata.normalize('NFC', value).strip()


def match_key(value):
    return ' '.join(text(value).split()).casefold()


def code(value, width, numeric_source=False):
    # Spreadsheet numeric cells can be rendered only at their declared exact width.
    # Never pad lost zeroes or infer a district code from a facility code.
    if numeric_source and type(value) in (int, float):
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value() or abs(number) >= 10 ** 15:
            raise ValueError('UNSAFE_NUMERIC_CODE')
        value = format(number, '.0f')
    value = text(value)
    if not re.fullmatch(r'[0-9]{' + str(width) + '}', value):
        raise ValueError('INVALID_CODE')
    return value


def coordinate(value, limit):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool) or not re.fullmatch(r'[+-]?\d+(?:\.\d+)?', str(value).strip()):
        raise ValueError('INVALID_COORDINATE')
    number = Decimal(str(value).strip())
    if not number.is_finite() or abs(number) > limit:
        raise ValueError('COORDINATE_OUT_OF_RANGE')
    return format(number, 'f')


def envelope(receipt, batch, number, locator, raw):
    return {'batch_id': batch, 'row_no': number, 'source_locator': locator,
            'source_id': receipt['source_id'], 'source_hash': receipt['sha256'],
            'source_edition': receipt['edition'], 'pipeline_version': VERSION,
            'raw_fields': raw, 'normalized_fields': {}, 'transform_log': [],
            'disposition': 'accepted', 'errors': [], 'data_quality_accepted': False}


def converted(row, field, before, after, reason, **details):
    row['normalized_fields'][field] = after
    if before != after or type(before) != type(after) or details:
        row['transform_log'].append({'field': field, 'before': before, 'after': after,
                                     'reason': reason, **details})


def reject(row, reason):
    row['errors'].append(reason)
    row['disposition'] = 'quarantined'


def normalize_facility(row, numeric_source=False, coordinates=False):
    raw = row['raw_fields']
    try:
        converted(row, 'facility_code', raw.get('kode_puskesmas'),
                  code(raw.get('kode_puskesmas'), 11, numeric_source), 'exact_declared_text_code')
        for field in FIELDS[2:]:
            value = text(raw.get(field))
            if not value:
                raise ValueError('MISSING_REQUIRED_FIELD')
            converted(row, field, raw[field], value, 'unicode_nfc_outer_trim')
        row['normalized_fields']['region_match_key'] = [match_key(raw['kabupaten']), match_key(raw['kecamatan'])]
        if coordinates:
            lat, lon = coordinate(raw.get('latitude'), 90), coordinate(raw.get('longitude'), 180)
            if (lat is None) != (lon is None):
                raise ValueError('INCOMPLETE_COORDINATE_PAIR')
            for field, value in [('latitude', lat), ('longitude', lon)]:
                converted(row, field, raw.get(field), value, 'declared_decimal_degrees', locale='en_US')
            row['normalized_fields'].update(location_status='missing' if lat is None else 'candidate_unverified',
                                             missing_reason='source_empty' if lat is None else None)
    except (ValueError, InvalidOperation) as exc:
        reject(row, str(exc))
    return row


def normalize_indicator(row, parsed, aggregate=False):
    try:
        if parsed.get('status') == 'quarantined':
            raise ValueError('PARSER_REJECTED: ' + parsed.get('reason', 'unknown'))
        year = parsed.get('observation_year')
        if type(year) is not int or not 1900 <= year <= 2100 or not parsed.get('year_evidence'):
            raise ValueError('UNKNOWN_PERIOD')
        if not parsed.get('unit_evidence'):
            raise ValueError('INVALID_UNIT')
        value = convert_value(parsed['raw_value'], parsed['indicator'], parsed['unit'], parsed['locale'])
        if parsed['indicator'] == 'area' and value is not None and Decimal(value) == 0:
            raise ValueError('NON_POSITIVE_AREA')
        for field, source in [('regency', 'regency_raw'), ('district', 'district_raw')]:
            converted(row, field, parsed[source], text(parsed[source]), 'unicode_nfc_outer_trim')
        converted(row, 'value', parsed['raw_value'], value, 'declared_locale_and_unit',
                  source_unit=parsed['unit'], target_unit='jiwa' if parsed['indicator'] == 'population' else 'km2',
                  locale=parsed['locale'], factor={'ribu jiwa': '1000', 'ha': '0.01', 'hektar': '0.01'}.get(parsed['unit'], '1'))
        row['normalized_fields'].update(metric=parsed['indicator'], observation_year=year,
                                         region_match_key=[match_key(parsed['regency_raw']), match_key(parsed['district_raw'])],
                                         missing_reason='source_missing_marker' if value is None else None,
                                         review_required=parsed.get('review_required', True))
        if aggregate:
            row['disposition'] = 'aggregate'
    except (KeyError, ValueError, InvalidOperation) as exc:
        reject(row, str(exc))
    return row


def read_table(data, spec):
    """Keep row width errors as rows; reject ambiguous headers/encoding as a batch."""
    expected = tuple(spec['headers'])
    if spec['format'] == 'csv':
        reader = csv.reader(io.StringIO(data.decode('utf-8-sig'), newline=''), strict=True)
        header = tuple(next(reader, []))
        rows = [(f'csv/row/{number}', values) for number, values in enumerate(reader, 2)]
        metadata = {'encoding': 'utf-8', 'bom': data.startswith(b'\xef\xbb\xbf'), 'header_rows': 1}
    else:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            sheet = book[spec['sheet']]
            iterator = sheet.iter_rows(values_only=True)
            for _ in range(spec.get('header_row', 1) - 1):
                next(iterator)
            header = tuple(next(iterator))
            rows = [(f'xlsx/{sheet.title}/row/{number}', list(values))
                    for number, values in enumerate(iterator, spec.get('header_row', 1) + 1)]
        finally:
            book.close()
        metadata = {'header_rows': spec.get('header_row', 1), 'encoding': 'xlsx_xml'}
    if len(set(header)) != len(header) or header != expected:
        raise ValueError('HEADER_MISMATCH')
    return rows, metadata


def normalize_table(data, spec, receipt, batch):
    source_rows, metadata = read_table(data, spec)
    output, blank_rows = [], 0
    for locator, values in source_rows:
        if all(value is None or value == '' for value in values):
            blank_rows += 1
            continue
        raw = dict(zip(spec['headers'], values))
        row = envelope(receipt, batch, len(output) + 1, locator, raw)
        row['raw_cells'] = values
        if len(values) != len(spec['headers']):
            reject(row, 'COLUMN_COUNT')
        elif spec['kind'] in ('facility', 'coordinates'):
            normalize_facility(row, spec['format'] == 'xlsx', spec['kind'] == 'coordinates')
        else:
            try:
                if spec['kind'] == 'combined':
                    regency, district = raw['Kabupaten'], raw['Kecamatan']
                    numeric = [('population', 'Jumlah Penduduk', 'jiwa'), ('area', 'Luas (km2)', 'km2'),
                               ('facility_count', 'Jumlah Puskesmas', 'jiwa')]
                else:
                    regency, district = raw['kabupaten'], raw['kecamatan']
                    numeric = [('facility_count', 'jumlah_puskesmas', 'jiwa')]
                aggregate = any(value in spec.get('total_labels', []) for value in (regency, district))
                if not aggregate and (not text(regency) or not text(district)):
                    raise ValueError('MISSING_REGION_NAME')
                row['normalized_fields'].update(regency=text(regency), district=text(district),
                                                 region_match_key=[match_key(regency), match_key(district)],
                                                 observation_year=None, period_status='unverified_historical_derivative')
                for target, field, unit in numeric:
                    value = convert_value(raw[field], 'area' if target == 'area' else 'population', unit, 'en_US')
                    if target == 'area' and value is not None and Decimal(value) == 0:
                        raise ValueError('NON_POSITIVE_AREA')
                    converted(row, target, raw[field], value, 'declared_derivative_numeric_column',
                              unit='facilities' if target == 'facility_count' else unit)
                if aggregate:
                    row['disposition'] = 'aggregate'
            except (KeyError, ValueError, InvalidOperation) as exc:
                reject(row, str(exc))
        output.append(row)
    metadata.update(blank_rows=blank_rows, parsed_data_rows=len(output))
    return output, metadata


def run(root, manifest, specs, output_root):
    root = Path(root).resolve()
    output_root = local_path(root, str(output_root), 'data/staging')
    if manifest.get('availability_complete') is not True:
        raise ValueError('SOURCE_INCOMPLETE')
    if not specs:
        raise ValueError('No source tables selected')
    by_hash = {}
    for receipt in manifest['files']:
        by_hash.setdefault(receipt['sha256'], []).append(receipt)
    if len({s['dataset'] for s in specs}) != len(specs):
        raise ValueError('Duplicate dataset')
    batch = digest({'pipeline_version': VERSION, 'specs': specs,
                    'source_plan': manifest['plan_sha256']})
    datasets, all_rows = [], []
    for spec in specs:
        if not re.fullmatch(r'[a-z0-9_-]+', spec['dataset']):
            raise ValueError('Invalid dataset identifier')
        receipts = [r for r in by_hash.get(spec['sha256'], []) if r['source_id'] == spec['source_id']]
        if len(receipts) != 1:
            raise ValueError('Source identity absent or ambiguous in frozen manifest')
        receipt = receipts[0]
        path = local_path(root, receipt['raw_path'], 'data/raw')
        data = checked_bytes(path, receipt)
        rows, metadata = [], {}
        if spec['format'] == 'facility_pdf':
            parsed = parse_pdf(path, {**receipt, 'edition': spec['edition']})
            if parsed['diagnostics']:
                raise ValueError('PDF_SCHEMA_DIAGNOSTICS')
            for group in ('rows', 'quarantined', 'aggregates'):
                for item in parsed[group]:
                    row = envelope(receipt, batch, len(rows) + 1, item['locator'], item['raw_fields'])
                    row['raw_cells'] = item['raw_cells']
                    if group == 'aggregates':
                        row['disposition'] = 'aggregate'
                    else:
                        normalize_facility(row)
                    if group == 'quarantined':
                        reject(row, 'PARSER_REJECTED')
                    rows.append(row)
            metadata = {'pages_scanned': parsed['pages_scanned'], 'detail_pages': parsed['detail_pages'],
                        'selection': 'Jawa Timur facility table rows'}
        elif spec['format'] == 'indicator_workbook':
            parsed = parse_workbook(path, receipt, spec['indicator'])
            if parsed['diagnostics']:
                raise ValueError('WORKBOOK_SCHEMA_DIAGNOSTICS')
            for group in ('rows', 'aggregates'):
                for item in parsed[group]:
                    raw = {k: v for k, v in item.items() if k not in ('value', 'status', 'acquired_at')}
                    row = envelope(receipt, batch, len(rows) + 1, item['locator'], raw)
                    normalize_indicator(row, item, group == 'aggregates')
                    rows.append(row)
            metadata = {'notes': parsed['notes'], 'source_kind': 'archived_derivative_workbook'}
        else:
            rows, metadata = normalize_table(data, spec, receipt, batch)
        for row in rows:
            row['dataset'] = spec['dataset']
        locators = [r['source_locator'] for r in rows]
        if len(set(locators)) != len(locators):
            raise ValueError('Duplicate source row locator')
        counts = Counter(r['disposition'] for r in rows)
        if 'expected_rows' in spec and len(rows) != spec['expected_rows']:
            raise ValueError('SOURCE_ROW_COUNT_MISMATCH: ' + spec['dataset'])
        datasets.append({'dataset': spec['dataset'], 'source_id': receipt['source_id'],
                         'source_hash': receipt['sha256'], 'raw': len(rows),
                         **{k: counts[k] for k in ('accepted', 'quarantined', 'aggregate')},
                         'metadata': metadata, 'content_sha256': digest(rows)})
        all_rows.extend(rows)
    # Complete all parsing before writing outputs. Publish a manifest last in a content-addressed batch.
    content = {'batch_id': batch, 'pipeline_version': VERSION, 'rows': all_rows}
    content_hash = digest(content)
    folder = output_root / batch / content_hash
    paths = {}
    for disposition, filename in [('accepted', 'staging.json'), ('quarantined', 'quarantine.json'), ('aggregate', 'aggregates.json')]:
        target = folder / filename
        payload = {**content, 'rows': [r for r in all_rows if r['disposition'] == disposition]}
        atomic_json(target, payload)
        paths[disposition] = {'path': target.relative_to(root).as_posix(),
                              'sha256': hashlib.sha256(target.read_bytes()).hexdigest(), 'rows': len(payload['rows'])}
    report = {'batch_id': batch, 'pipeline_version': VERSION, 'content_sha256': content_hash,
              'datasets': datasets, 'outputs': paths, 'data_quality_accepted': False,
              'source_plan_sha256': manifest['plan_sha256'], 'raw': len(all_rows),
              'counts': {k: v['rows'] for k, v in paths.items()}}
    atomic_json(folder / 'manifest.json', report)
    return report
